from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

import pandas as pd


PROJECT_ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = PROJECT_ROOT / "data" / "raw"
REPORT_DIR = PROJECT_ROOT / "reports"

DEFAULT_ZIP = RAW_DIR / "ODBus_2023.zip"


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def decode_bytes(raw: bytes) -> tuple[str, str]:
    """Try common encodings used by government/open-data files."""
    for encoding in (
        "utf-8-sig",
        "utf-8",
        "cp1252",
        "latin-1",
    ):
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            pass

    return raw.decode("latin-1", errors="replace"), "latin-1-replace"


def read_csv_from_zip(
    zf: zipfile.ZipFile,
    member: str,
) -> tuple[pd.DataFrame, str]:
    """Read a CSV directly from the ZIP."""
    with zf.open(member) as f:
        raw = f.read()

    text, encoding = decode_bytes(raw)

    df = pd.read_csv(
        io.StringIO(text),
        low_memory=False,
    )

    return df, encoding


def clean(value) -> str:
    if pd.isna(value):
        return ""

    return str(value).strip()


def markdown_table(
    rows: list[list],
    headers: list[str],
) -> str:

    lines = []

    lines.append(
        "| " + " | ".join(clean(x) for x in headers) + " |"
    )

    lines.append(
        "| " + " | ".join("---" for _ in headers) + " |"
    )

    for row in rows:
        values = []

        for value in row:
            value = clean(value)
            value = value.replace("|", "\\|")
            value = value.replace("\n", " ")

            values.append(value)

        lines.append(
            "| " + " | ".join(values) + " |"
        )

    return "\n".join(lines)


def find_member(
    members: list[str],
    filename: str,
) -> str | None:

    filename_lower = filename.lower()

    for member in members:
        if member.lower().endswith(filename_lower):
            return member

    return None


def find_column(
    columns,
    candidates: list[str],
) -> str | None:

    normalized = {
        str(column).strip().lower(): column
        for column in columns
    }

    for candidate in candidates:
        candidate_lower = candidate.lower()

        if candidate_lower in normalized:
            return normalized[candidate_lower]

    return None


# ---------------------------------------------------------------------------
# Metadata PDF/DOCX extraction
# ---------------------------------------------------------------------------

def extract_pdf_text(zf: zipfile.ZipFile, member: str) -> str:
    """
    Extract text from the PDF.

    Requires pypdf:
        python -m pip install pypdf
    """
    try:
        from pypdf import PdfReader
    except ImportError:
        return (
            "PDF TEXT EXTRACTION UNAVAILABLE.\n\n"
            "Install pypdf with:\n"
            "python -m pip install pypdf\n"
        )

    raw = zf.read(member)

    reader = PdfReader(io.BytesIO(raw))

    pages = []

    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            text = (
                f"[Could not extract page {page_number}: {exc}]"
            )

        pages.append(
            f"\n--- PDF PAGE {page_number} ---\n{text}"
        )

    return "\n".join(pages)


def extract_docx_text(zf: zipfile.ZipFile, member: str) -> str:
    """
    Extract DOCX text.

    Requires python-docx:
        python -m pip install python-docx
    """
    try:
        from docx import Document
    except ImportError:
        return (
            "DOCX TEXT EXTRACTION UNAVAILABLE.\n\n"
            "Install python-docx with:\n"
            "python -m pip install python-docx\n"
        )

    raw = zf.read(member)

    document = Document(io.BytesIO(raw))

    paragraphs = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    # Also capture table contents because metadata often lives in tables.
    tables = []

    for table_index, table in enumerate(document.tables, start=1):

        tables.append(
            f"\n--- DOCX TABLE {table_index} ---"
        )

        for row in table.rows:
            values = [
                cell.text.strip()
                for cell in row.cells
            ]

            tables.append(
                " | ".join(values)
            )

    return "\n".join(
        paragraphs + tables
    )


# ---------------------------------------------------------------------------
# Metadata keyword extraction
# ---------------------------------------------------------------------------

def extract_relevant_sections(
    text: str,
) -> dict[str, list[str]]:

    keywords = {
        "date": [
            "date",
            "year",
            "reference period",
            "reference date",
            "updated",
            "update",
            "vintage",
        ],
        "license": [
            "licence",
            "license",
            "copyright",
            "open government",
            "terms",
        ],
        "coverage": [
            "coverage",
            "province",
            "territory",
            "geographic",
            "geography",
        ],
        "business_unit": [
            "business",
            "establishment",
            "enterprise",
            "licence",
            "license",
            "record",
        ],
        "employee": [
            "employee",
            "employment",
            "worker",
            "staff",
        ],
        "provider": [
            "provider",
            "source",
            "dataset",
            "contributor",
        ],
        "identifier": [
            "business_id",
            "identifier",
            "id number",
            "unique",
            "key",
        ],
    }

    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    result = {}

    for category, words in keywords.items():

        matches = []

        for index, line in enumerate(lines):

            lower = line.lower()

            if any(
                word.lower() in lower
                for word in words
            ):

                # Capture nearby lines for context.
                start = max(0, index - 1)
                end = min(
                    len(lines),
                    index + 3,
                )

                block = "\n".join(
                    lines[start:end]
                )

                if block not in matches:
                    matches.append(block)

        result[category] = matches[:30]

    return result


# ---------------------------------------------------------------------------
# Provider analysis
# ---------------------------------------------------------------------------

def analyze_provider_data(
    df: pd.DataFrame,
) -> dict:

    result = {
        "provider_column": None,
        "province_column": None,
        "provider_count": None,
        "provider_rows": [],
        "provider_province_rows": [],
    }

    provider_col = find_column(
        df.columns,
        [
            "provider",
        ],
    )

    province_col = find_column(
        df.columns,
        [
            "prov_terr",
            "province",
            "province_territory",
        ],
    )

    result["provider_column"] = provider_col
    result["province_column"] = province_col

    if not provider_col:
        return result

    provider_series = (
        df[provider_col]
        .fillna("<MISSING>")
        .astype(str)
        .str.strip()
    )

    provider_counts = (
        provider_series
        .value_counts()
    )

    result["provider_count"] = len(
        provider_counts
    )

    result["provider_rows"] = [
        [
            provider,
            int(count),
            round(
                count / len(df) * 100,
                2,
            ),
        ]
        for provider, count
        in provider_counts.items()
    ]

    if province_col:

        temp = pd.DataFrame(
            {
                "provider": provider_series,
                "province": (
                    df[province_col]
                    .fillna("<MISSING>")
                    .astype(str)
                    .str.strip()
                ),
            }
        )

        cross = (
            temp.groupby(
                ["provider", "province"]
            )
            .size()
            .reset_index(
                name="records"
            )
            .sort_values(
                "records",
                ascending=False,
            )
        )

        result["provider_province_rows"] = [
            [
                row["provider"],
                row["province"],
                int(row["records"]),
            ]
            for _, row in cross.iterrows()
        ]

    return result


# ---------------------------------------------------------------------------
# Main report
# ---------------------------------------------------------------------------

def build_report(
    zip_path: Path,
    layout_df: pd.DataFrame,
    sources_df: pd.DataFrame,
    main_df: pd.DataFrame,
    layout_encoding: str,
    sources_encoding: str,
    main_encoding: str,
    pdf_text: str,
    docx_text: str,
) -> str:

    lines = []

    lines.append("# ODBus Metadata & Source Inspection")
    lines.append("")
    lines.append(
        "Generated from `ODBus_2023.zip`."
    )
    lines.append("")

    # -----------------------------------------------------------------------
    # Overview
    # -----------------------------------------------------------------------

    lines.append("## 1. Dataset overview")
    lines.append("")

    lines.append(
        f"- Archive: `{zip_path.name}`"
    )
    lines.append(
        f"- Main dataset rows: **{len(main_df):,}**"
    )
    lines.append(
        f"- Main dataset columns: **{len(main_df.columns):,}**"
    )
    lines.append(
        f"- Main dataset encoding: `{main_encoding}`"
    )
    lines.append(
        f"- Record-layout encoding: `{layout_encoding}`"
    )
    lines.append(
        f"- Sources-table encoding: `{sources_encoding}`"
    )
    lines.append("")

    # -----------------------------------------------------------------------
    # Record layout
    # -----------------------------------------------------------------------

    lines.append("## 2. ODBus record layout")
    lines.append("")

    lines.append(
        "This section reproduces the structured field-definition table "
        "contained in `ODBus-record-layout.csv`."
    )
    lines.append("")

    lines.append(
        markdown_table(
            layout_df.astype(str).values.tolist(),
            [str(c) for c in layout_df.columns],
        )
    )

    lines.append("")

    # -----------------------------------------------------------------------
    # Source table
    # -----------------------------------------------------------------------

    lines.append("## 3. ODBus source table")
    lines.append("")

    lines.append(
        f"`ODBus_Sources.csv` contains "
        f"**{len(sources_df):,} rows** and "
        f"**{len(sources_df.columns):,} columns**."
    )

    lines.append("")

    lines.append(
        markdown_table(
            sources_df.astype(str).head(100).values.tolist(),
            [str(c) for c in sources_df.columns],
        )
    )

    lines.append("")

    if len(sources_df) > 100:
        lines.append(
            f"> Only the first 100 of {len(sources_df):,} "
            "source-table rows are displayed."
        )
        lines.append("")

    # -----------------------------------------------------------------------
    # Provider analysis
    # -----------------------------------------------------------------------

    provider_analysis = analyze_provider_data(
        main_df
    )

    lines.append("## 4. Provider analysis")
    lines.append("")

    if provider_analysis["provider_column"]:

        lines.append(
            f"- Provider column: "
            f"`{provider_analysis['provider_column']}`"
        )

        if provider_analysis["province_column"]:
            lines.append(
                f"- Province column: "
                f"`{provider_analysis['province_column']}`"
            )

        lines.append(
            f"- Distinct provider values: "
            f"**{provider_analysis['provider_count']:,}**"
        )

        lines.append("")

        lines.append("### Records by provider")
        lines.append("")

        lines.append(
            markdown_table(
                provider_analysis["provider_rows"],
                [
                    "Provider",
                    "Records",
                    "Share %",
                ],
            )
        )

        lines.append("")

        if provider_analysis[
            "provider_province_rows"
        ]:

            lines.append(
                "### Provider × province distribution"
            )
            lines.append("")

            lines.append(
                markdown_table(
                    provider_analysis[
                        "provider_province_rows"
                    ],
                    [
                        "Provider",
                        "Province/Territory",
                        "Records",
                    ],
                )
            )

            lines.append("")

    else:
        lines.append(
            "No `provider` column was detected in the main dataset."
        )
        lines.append("")

    # -----------------------------------------------------------------------
    # Main field list
    # -----------------------------------------------------------------------

    lines.append("## 5. Main dataset fields")
    lines.append("")

    for column in main_df.columns:
        lines.append(f"- `{column}`")

    lines.append("")

    # -----------------------------------------------------------------------
    # Metadata text
    # -----------------------------------------------------------------------

    lines.append("## 6. PDF metadata text")
    lines.append("")

    if pdf_text.strip():
        lines.append(
            "The following is extracted text from the supplied "
            "`ODBus Metadata.pdf`. It is included for inspection; "
            "interpretation should follow the original metadata."
        )
        lines.append("")
        lines.append("```text")
        lines.append(pdf_text[:100_000])
        lines.append("```")
    else:
        lines.append(
            "No PDF text was extracted."
        )

    lines.append("")

    lines.append("## 7. DOCX metadata text")
    lines.append("")

    if docx_text.strip():
        lines.append(
            "The following is extracted text from the supplied "
            "`ODBus Metadata.docx`."
        )
        lines.append("")
        lines.append("```text")
        lines.append(docx_text[:100_000])
        lines.append("```")
    else:
        lines.append(
            "No DOCX text was extracted."
        )

    lines.append("")

    # -----------------------------------------------------------------------
    # Keyword findings
    # -----------------------------------------------------------------------

    combined_metadata = (
        pdf_text
        + "\n"
        + docx_text
    )

    relevant = extract_relevant_sections(
        combined_metadata
    )

    lines.append(
        "## 8. Metadata keyword findings"
    )
    lines.append("")

    for category, matches in relevant.items():

        lines.append(
            f"### {category.replace('_', ' ').title()}"
        )
        lines.append("")

        if matches:

            for match in matches:
                lines.append(
                    f"```text\n{match}\n```"
                )

        else:
            lines.append(
                "_No matching metadata text detected._"
            )

        lines.append("")

    # -----------------------------------------------------------------------
    # Investigation questions
    # -----------------------------------------------------------------------

    lines.append(
        "## 9. Questions to resolve from this inspection"
    )
    lines.append("")

    questions = [
        "What exactly does one ODBus row represent?",
        "What does `business_id_no` uniquely identify?",
        "Is `business_id_no` unique globally or only within a provider/source?",
        "What does `provider` mean?",
        "What organization/dataset does each provider value represent?",
        "Why is geographic coverage concentrated in certain provinces/territories?",
        "Which source contributes employee information?",
        "How is `total_no_employees` constructed?",
        "What does `..` mean in each important field?",
        "What is the reference date or collection period?",
        "What is the update frequency?",
        "What are the licensing/commercial-use conditions?",
        "Does the source support historical/change detection?",
        "Which fields are source-derived versus ODBus-derived?",
        "Can records represent multiple establishments or licences belonging to one business?",
    ]

    for index, question in enumerate(
        questions,
        start=1,
    ):
        lines.append(
            f"{index}. {question}"
        )

    lines.append("")

    # -----------------------------------------------------------------------
    # Important warning
    # -----------------------------------------------------------------------

    lines.append(
        "## 10. Interpretation status"
    )
    lines.append("")

    lines.append(
        "> This report is an evidence-gathering artifact. "
        "It does not automatically resolve the semantic meaning of "
        "the ODBus fields. Decisions about record grain, identifiers, "
        "coverage, employee classification, licensing, or production "
        "use should be based on the supplied metadata and source "
        "documentation."
    )

    lines.append("")

    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():

    if not DEFAULT_ZIP.exists():
        raise FileNotFoundError(
            f"Could not find: {DEFAULT_ZIP}"
        )

    REPORT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    print("=" * 70)
    print("ODBus Metadata & Source Inspection")
    print("=" * 70)

    print(f"\nOpening: {DEFAULT_ZIP}")

    with zipfile.ZipFile(
        DEFAULT_ZIP,
        "r",
    ) as zf:

        members = zf.namelist()

        layout_member = find_member(
            members,
            "ODBus-record-layout.csv",
        )

        sources_member = find_member(
            members,
            "ODBus_Sources.csv",
        )

        main_member = find_member(
            members,
            "ODBus_v1.csv",
        )

        pdf_member = find_member(
            members,
            "ODBus Metadata.pdf",
        )

        docx_member = find_member(
            members,
            "ODBus Metadata.docx",
        )

        required = {
            "record layout": layout_member,
            "sources": sources_member,
            "main dataset": main_member,
        }

        for name, member in required.items():

            if not member:
                raise RuntimeError(
                    f"Could not find required {name} file."
                )

        print("\nReading record layout...")
        layout_df, layout_encoding = (
            read_csv_from_zip(
                zf,
                layout_member,
            )
        )

        print(
            f"  Rows: {len(layout_df):,}"
        )
        print(
            f"  Columns: {len(layout_df.columns):,}"
        )
        print(
            f"  Encoding: {layout_encoding}"
        )

        print("\nReading source table...")
        sources_df, sources_encoding = (
            read_csv_from_zip(
                zf,
                sources_member,
            )
        )

        print(
            f"  Rows: {len(sources_df):,}"
        )
        print(
            f"  Columns: {len(sources_df.columns):,}"
        )
        print(
            f"  Encoding: {sources_encoding}"
        )

        print("\nReading main dataset...")
        main_df, main_encoding = (
            read_csv_from_zip(
                zf,
                main_member,
            )
        )

        print(
            f"  Rows: {len(main_df):,}"
        )
        print(
            f"  Columns: {len(main_df.columns):,}"
        )
        print(
            f"  Encoding: {main_encoding}"
        )

        # -------------------------------------------------------------------
        # PDF
        # -------------------------------------------------------------------

        if pdf_member:

            print("\nExtracting PDF metadata...")
            pdf_text = extract_pdf_text(
                zf,
                pdf_member,
            )

            print(
                f"  Extracted approximately "
                f"{len(pdf_text):,} characters"
            )

        else:

            print("\nPDF metadata not found.")
            pdf_text = ""

        # -------------------------------------------------------------------
        # DOCX
        # -------------------------------------------------------------------

        if docx_member:

            print("\nExtracting DOCX metadata...")
            docx_text = extract_docx_text(
                zf,
                docx_member,
            )

            print(
                f"  Extracted approximately "
                f"{len(docx_text):,} characters"
            )

        else:

            print("\nDOCX metadata not found.")
            docx_text = ""

    # -----------------------------------------------------------------------
    # Build report
    # -----------------------------------------------------------------------

    report = build_report(
        zip_path=DEFAULT_ZIP,
        layout_df=layout_df,
        sources_df=sources_df,
        main_df=main_df,
        layout_encoding=layout_encoding,
        sources_encoding=sources_encoding,
        main_encoding=main_encoding,
        pdf_text=pdf_text,
        docx_text=docx_text,
    )

    output_path = (
        REPORT_DIR
        / "ODBUS_METADATA_INSPECTION.md"
    )

    output_path.write_text(
        report,
        encoding="utf-8",
    )

    print("\n" + "=" * 70)
    print("Inspection complete")
    print("=" * 70)

    print(
        f"\nReport written to:\n"
        f"  {output_path}"
    )

    print("\nNext step:")
    print(
        "Review reports\\ODBUS_METADATA_INSPECTION.md"
    )


if __name__ == "__main__":
    main()